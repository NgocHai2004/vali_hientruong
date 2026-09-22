import io
import re
from datetime import datetime
from typing import Optional, List
from fastapi import APIRouter, Depends, HTTPException, Query, Request, UploadFile, File
from fastapi.responses import StreamingResponse
import httpx
from openpyxl import Workbook, load_workbook

from core.config import (
    FP_SERVICE_URL,
    FP_MATCH_FINGER,
    FP_LEFT_THUMB_THRESHOLD,
    FP_SINGLE_THRESHOLD,
    FP_REQUIRED_FINGER_COUNT,
    FP_FINGER_CODES,
    FP_KEY_BY_CODE,
)
from core.security import get_current_user, scope_filter
from schemas.detainees import (
    DetaineeIn,
    TransferBody,
    MatchFingerprintReq,
    MatchFingerprintSingleReq,
)
from services.hbie_matcher import _spawn_case_rematch
import db.mongo as db_module

router = APIRouter(tags=["Detainees"])

_MATCH_PROJECTION = {
    "personal_id": 1, "full_name": 1, "cccd_number": 1, "gender": 1, "dob": 1,
    "cell_code": 1, "custody_type": 1, "facility_code": 1, "sub_camp_code": 1,
    "charge": 1, "hometown": 1, "address": 1,
    "photos.portrait_front": 1, "photos.cccd_front": 1,
    "created_at": 1, "created_by": 1,
}

EXCEL_COLS = [
    ("personal_id", "Số định danh"),
    ("full_name", "Họ và tên"),
    ("gender", "Giới tính"),
    ("dob", "Ngày sinh"),
    ("cccd_number", "Số CCCD"),
    ("hometown", "Quê quán"),
    ("address", "Địa chỉ"),
    ("ethnicity", "Dân tộc"),
    ("religion", "Tôn giáo"),
    ("facility_code", "Nơi giam giữ"),
    ("sub_camp_code", "Phân trại"),
    ("cell_code", "Mã buồng"),
    ("charge", "Tội danh"),
    ("date_in", "Ngày vào"),
    ("note", "Ghi chú"),
]


def _require_capture_fields(body: DetaineeIn) -> None:
    missing = []
    if not body.cccd_number:
        missing.append("Số CCCD (12 chữ số)")
    if missing:
        raise HTTPException(400, "Thiếu thông tin bắt buộc: " + ", ".join(missing))


async def _find_duplicates(full_name: str, dob: Optional[str], gender: str, exclude_id: Optional[str] = None) -> List[dict]:
    if not full_name:
        return []
    q = {"full_name": full_name.strip(), "gender": gender}
    if dob:
        q["dob"] = dob
    if exclude_id:
        q["_id"] = {"$ne": db_module._oid(exclude_id)}
    return [db_module._s(d) async for d in db_module.db.detainees.find(q).limit(5)]


def _ensure_can_touch(doc: dict, user: dict) -> None:
    if user.get("role") == "admin":
        return
    if doc.get("created_by") != user["username"]:
        raise HTTPException(403, "Bạn chỉ được thao tác trên hồ sơ do chính mình đăng ký")


def _has_roll_photos(photos: Optional[dict]) -> bool:
    ph = photos or {}
    return any(ph.get(k) for k in FP_KEY_BY_CODE.values())


def _changed_roll_fingers(old: Optional[dict], new: Optional[dict]) -> list:
    o, n = old or {}, new or {}
    return [c for c, k in FP_KEY_BY_CODE.items() if (o.get(k) or "") != (n.get(k) or "")]


@router.get("/api/detainees/export/xlsx")
async def export_xlsx(user: dict = Depends(get_current_user)):
    wb = Workbook()
    ws = wb.active
    ws.title = "Nghi pham"
    ws.append([h for _, h in EXCEL_COLS])
    async for d in db_module.db.detainees.find({}).sort("personal_id", 1):
        row = []
        for k, _ in EXCEL_COLS:
            v = d.get(k, "")
            row.append(v if v is not None else "")
        ws.append(row)
    for col in ws.columns:
        letter = col[0].column_letter
        ws.column_dimensions[letter].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"nghi_pham_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


@router.get("/api/detainees/template/xlsx")
async def template_xlsx(user: dict = Depends(get_current_user)):
    wb = Workbook()
    ws = wb.active
    ws.title = "Mau nhap"
    ws.append([h for _, h in EXCEL_COLS])
    ws.append(["", "Nguyễn Văn Mẫu", "male", "15/03/1990", "001090123456", "Hà Nội", "Số 1, Hà Nội", "Kinh", "Không", "A01", "Trộm cắp tài sản", "01/01/2026", ""])
    for col in ws.columns:
        letter = col[0].column_letter
        ws.column_dimensions[letter].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="mau_import.xlsx"'},
    )


@router.post("/api/detainees/import/xlsx")
async def import_xlsx(file: UploadFile = File(...), request: Request = None, user: dict = Depends(get_current_user)):
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(400, "Chỉ nhận file .xlsx")
    data = await file.read()
    wb = load_workbook(io.BytesIO(data), read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(400, "File rỗng")
    header = [str(x or "").strip() for x in rows[0]]
    header_map = {h: idx for idx, h in enumerate(header)}
    inserted, errors = 0, []
    now = datetime.utcnow()
    for i, r in enumerate(rows[1:], start=2):
        if not r or not any(r):
            continue
        try:
            def get(col_key):
                _, label = next((c for c in EXCEL_COLS if c[0] == col_key), (None, None))
                idx = header_map.get(label) if label else None
                if idx is None or idx >= len(r):
                    return None
                v = r[idx]
                return str(v).strip() if v is not None else None
            full_name = get("full_name")
            if not full_name:
                errors.append(f"Dòng {i}: thiếu Họ và tên")
                continue
            dob = db_module._parse_dob(get("dob"))
            date_in = db_module._parse_dob(get("date_in"))
            personal_id = (get("personal_id") or "").strip()
            if not personal_id:
                errors.append(f"Dòng {i}: thiếu mã nghi phạm (personal_id)")
                continue
            doc = {
                "personal_id": personal_id,
                "full_name": full_name,
                "gender": (get("gender") or "male").lower(),
                "dob": dob,
                "cccd_number": get("cccd_number") or "",
                "hometown": get("hometown"),
                "address": get("address"),
                "ethnicity": get("ethnicity"),
                "religion": get("religion"),
                "cell_code": get("cell_code"),
                "charge": get("charge"),
                "date_in": date_in,
                "note": get("note"),
                "created_at": now,
                "updated_at": now,
                "created_by": user["username"],
            }
            try:
                await db_module.db.detainees.insert_one(doc)
                inserted += 1
            except Exception as e:
                if "duplicate key" in str(e):
                    errors.append(f"Dòng {i}: số định danh {personal_id} đã tồn tại")
                else:
                    errors.append(f"Dòng {i}: {e}")
        except Exception as e:
            errors.append(f"Dòng {i}: {e}")
    await db_module._log(request, user, "import", "detainee", "", {"inserted": inserted, "errors": len(errors)})
    return {"inserted": inserted, "errors": errors}


@router.get("/api/detainees")
async def list_detainees(
    q: str = Query("", alias="q"),
    cell_code: str = Query(""),
    gender: str = Query(""),
    case_id: Optional[str] = Query(None),
    skip: int = Query(0, ge=0),
    limit: int = Query(20, ge=1, le=200),
    user: dict = Depends(get_current_user),
):
    filt = scope_filter(user)
    if case_id:
        c_oid = db_module._oid(case_id)
        if c_oid:
            filt["case_id"] = {"$in": [c_oid, case_id, str(c_oid)]}
        else:
            filt["case_id"] = case_id
    if q:
        rx = re.escape(q.strip())
        filt["$or"] = [
            {"full_name": {"$regex": rx, "$options": "i"}},
            {"cccd_number": {"$regex": rx, "$options": "i"}},
            {"personal_id": {"$regex": rx, "$options": "i"}},
        ]
    if cell_code:
        filt["cell_code"] = cell_code
    if gender:
        filt["gender"] = gender
    total = await db_module.db.detainees.count_documents(filt)
    items = [
        db_module._s(d)
        async for d in db_module.db.detainees.find(filt).sort("created_at", -1).skip(skip).limit(limit)
    ]
    return {"total": total, "items": items, "skip": skip, "limit": limit}


@router.get("/api/detainees/check-cccd")
async def check_cccd(
    cccd_number: str = Query("", min_length=1),
    user: dict = Depends(get_current_user),
):
    cccd = re.sub(r"\D", "", cccd_number or "")
    if not cccd:
        raise HTTPException(400, "Thiếu số CCCD.")
    doc = await db_module.db.detainees.find_one(
        {"cccd_number": cccd},
        _MATCH_PROJECTION,
    )
    return {"matched": doc is not None, "detainee": db_module._s(doc) if doc else None}


@router.post("/api/detainees/check-duplicate")
async def check_duplicate(body: DetaineeIn, user: dict = Depends(get_current_user)):
    dob = db_module._parse_dob(body.dob)
    dups = await _find_duplicates(body.full_name, dob, body.gender)
    return {"count": len(dups), "duplicates": dups}


@router.get("/api/detainees/by-personal-id/{personal_id}")
async def get_detainee_by_personal_id(personal_id: str, user: dict = Depends(get_current_user)):
    doc = await db_module.db.detainees.find_one({"personal_id": personal_id})
    if not doc:
        raise HTTPException(404, "Không tìm thấy hồ sơ")
    _ensure_can_touch(doc, user)
    return db_module._s(doc)


@router.get("/api/detainees/{det_id}")
async def get_detainee(det_id: str, user: dict = Depends(get_current_user)):
    doc = await db_module.db.detainees.find_one({"_id": db_module._oid(det_id)})
    if not doc:
        raise HTTPException(404, "Không tìm thấy hồ sơ")
    _ensure_can_touch(doc, user)
    return db_module._s(doc)


@router.post("/api/detainees/match_fingerprint")
async def match_fingerprint(body: MatchFingerprintReq, user: dict = Depends(get_current_user)):
    fingers = body.fingers or {}
    present = [c for c in FP_FINGER_CODES if (fingers.get(c) or "").strip()]
    if len(present) < FP_REQUIRED_FINGER_COUNT:
        raise HTTPException(
            400,
            f"Phai quet du {FP_REQUIRED_FINGER_COUNT} ngon moi tra cuu (hien co {len(present)}).",
        )
    query_tmpl = (fingers.get(FP_MATCH_FINGER) or "").strip()
    if not query_tmpl:
        raise HTTPException(400, f"Thieu template cua ngon {FP_MATCH_FINGER}.")

    cursor = db_module.db.detainees.find(
        {"photos.fp_templates": {"$exists": True, "$ne": {}}},
        {
            "personal_id": 1, "full_name": 1, "cccd_number": 1, "gender": 1, "dob": 1,
            "cell_code": 1, "charge": 1, "hometown": 1, "address": 1,
            "photos.fp_templates": 1, "photos.portrait_front": 1, "photos.cccd_front": 1,
            "created_at": 1,
        },
    )

    matches: list[dict] = []
    async with httpx.AsyncClient(timeout=10.0) as client:
        async for det in cursor:
            fp_templates = (det.get("photos") or {}).get("fp_templates") or {}
            stored_tmpl = fp_templates.get(FP_MATCH_FINGER)
            if not stored_tmpl:
                continue
            try:
                resp = await client.post(
                    f"{FP_SERVICE_URL}/api/match_pair",
                    json={"t1_b64": query_tmpl, "t2_b64": stored_tmpl},
                )
                if resp.status_code != 200:
                    continue
                score = int(resp.json().get("score", 0))
            except Exception:
                continue
            if score > FP_LEFT_THUMB_THRESHOLD:
                matches.append({
                    "detainee": db_module._s(det),
                    "score": score,
                    "finger_code": FP_MATCH_FINGER,
                })

    matches.sort(key=lambda m: -m["score"])
    top = matches[:10]

    return {
        "matched": len(top) > 0,
        "total": len(top),
        "items": [
            {**m["detainee"], "match_score": m["score"], "match_finger": m["finger_code"]}
            for m in top
        ],
        "score": top[0]["score"] / 100.0 if top else 0.0,
    }


@router.post("/api/detainees/match_fingerprint_single")
async def match_fingerprint_single(body: MatchFingerprintSingleReq, user: dict = Depends(get_current_user)):
    if not body.template_b64:
        raise HTTPException(400, "Thieu template van tay.")

    cursor = db_module.db.detainees.find(
        {"photos.fp_templates": {"$exists": True, "$ne": {}}},
        {
            "personal_id": 1, "full_name": 1, "cccd_number": 1, "gender": 1, "dob": 1,
            "cell_code": 1, "charge": 1, "hometown": 1, "address": 1,
            "photos.fp_templates": 1, "photos.portrait_front": 1, "photos.cccd_front": 1,
            "created_at": 1,
        },
    )

    matches: list[dict] = []
    async with httpx.AsyncClient(timeout=10.0) as client:
        async for det in cursor:
            fp_templates = (det.get("photos") or {}).get("fp_templates") or {}
            best_score = 0
            best_finger = None
            for code, tmpl_b64 in fp_templates.items():
                if not tmpl_b64:
                    continue
                try:
                    resp = await client.post(
                        f"{FP_SERVICE_URL}/api/match_pair",
                        json={"t1_b64": body.template_b64, "t2_b64": tmpl_b64},
                    )
                    if resp.status_code != 200:
                        continue
                    score = int(resp.json().get("score", 0))
                except Exception:
                    continue
                if score > best_score:
                    best_score = score
                    best_finger = code
            if best_score > FP_SINGLE_THRESHOLD:
                matches.append({
                    "detainee": db_module._s(det),
                    "score": best_score,
                    "finger_code": best_finger,
                })

    matches.sort(key=lambda m: -m["score"])
    top = matches[:10]

    return {
        "matched": len(top) > 0,
        "total": len(top),
        "items": [
            {**m["detainee"], "match_score": m["score"], "match_finger": m["finger_code"]}
            for m in top
        ],
        "score": top[0]["score"] / 100.0 if top else 0.0,
    }


@router.post("/api/detainees")
async def create_detainee(body: DetaineeIn, request: Request, user: dict = Depends(get_current_user)):
    if user.get("role") == "admin":
        raise HTTPException(403, "Tài khoản quản trị hệ thống không thu nhận hồ sơ. Việc này do cán bộ thu nhận thực hiện.")
    if not body.case_id:
        raise HTTPException(400, "Bạn phải chọn vụ án trước khi tạo hồ sơ.")
    case_doc = await db_module.db.cases.find_one({"_id": db_module._oid(body.case_id)})
    if not case_doc:
        raise HTTPException(400, "Vụ án không tồn tại.")
    db_module._ensure_case_editable(case_doc)
    _require_capture_fields(body)
    dob = db_module._parse_dob(body.dob)
    now = datetime.utcnow()

    personal_id = (body.personal_id or "").strip()
    if not personal_id:
        raise HTTPException(400, "Thiếu mã nghi phạm (personal_id).")
    if await db_module.db.detainees.find_one({"personal_id": personal_id}):
        raise HTTPException(400, f"Mã nghi phạm '{personal_id}' đã có trong hệ thống.")

    doc = body.model_dump()
    doc.pop("case_id", None)
    doc.update({
        "personal_id": personal_id,
        "cccd_number": body.cccd_number or "",
        "dob": dob,
        "date_in": db_module._parse_dob(body.date_in),
        "issued_date": db_module._parse_dob(body.issued_date),
        "expiry_date": db_module._parse_dob(body.expiry_date),
        "created_at": now,
        "updated_at": now,
        "created_by": user["username"],
        "case_id": case_doc["_id"],
    })
    try:
        res = await db_module.db.detainees.insert_one(doc)
    except Exception as e:
        if "duplicate key" in str(e):
            raise HTTPException(400, f"Số định danh '{personal_id}' đã tồn tại (đồng thời), vui lòng thử lại.")
        raise
    doc["_id"] = res.inserted_id
    await db_module.db.cases.update_one({"_id": case_doc["_id"]}, {"$set": {"updated_at": now}})
    await db_module._log(request, user, "create", "detainee", personal_id, {"full_name": body.full_name, "case": case_doc.get("code")}, ref_id=str(res.inserted_id), case_id=case_doc["_id"])
    if _has_roll_photos(doc.get("photos")):
        _spawn_case_rematch(case_doc)
    return db_module._s(doc)


@router.patch("/api/detainees/{det_id}")
async def update_detainee(det_id: str, body: DetaineeIn, request: Request, user: dict = Depends(get_current_user)):
    existing = await db_module.db.detainees.find_one({"_id": db_module._oid(det_id)})
    if not existing:
        raise HTTPException(404, "Không tìm thấy hồ sơ")
    _ensure_can_touch(existing, user)
    cid = existing.get("case_id")
    if cid is not None:
        case_doc = await db_module.db.cases.find_one({"_id": cid})
        if case_doc:
            db_module._ensure_case_editable(case_doc)
    upd = body.model_dump()
    upd.pop("case_id", None)
    upd["dob"] = db_module._parse_dob(body.dob)
    upd["date_in"] = db_module._parse_dob(body.date_in)
    upd["issued_date"] = db_module._parse_dob(body.issued_date)
    upd["expiry_date"] = db_module._parse_dob(body.expiry_date)
    new_pid = (body.personal_id or "").strip()
    if new_pid:
        conflict = await db_module.db.detainees.find_one({"personal_id": new_pid, "_id": {"$ne": db_module._oid(det_id)}})
        if conflict:
            raise HTTPException(400, f"Mã nghi phạm '{new_pid}' đã có trong hồ sơ khác.")
        upd["personal_id"] = new_pid
        upd["cccd_number"] = body.cccd_number or upd.get("cccd_number", "")
    upd["updated_at"] = datetime.utcnow()
    changed_fp = _changed_roll_fingers(existing.get("photos"), upd.get("photos"))
    doc = await db_module.db.detainees.find_one_and_update({"_id": db_module._oid(det_id)}, {"$set": upd}, return_document=True)
    if changed_fp:
        await db_module.db.detainees.update_one(
            {"_id": db_module._oid(det_id)},
            {"$unset": {f"fp_features.{c}": "" for c in changed_fp}},
        )
        for c in changed_fp:
            (doc.get("fp_features") or {}).pop(c, None)
    await db_module._log(request, user, "update", "detainee", doc.get("personal_id", det_id), {"full_name": body.full_name}, ref_id=det_id, case_id=doc.get("case_id"))
    if changed_fp and _has_roll_photos(doc.get("photos")):
        cd = await db_module.db.cases.find_one({"_id": doc.get("case_id")}) if doc.get("case_id") else None
        if cd:
            _spawn_case_rematch(cd)
    return db_module._s(doc)


@router.delete("/api/detainees/{det_id}")
async def delete_detainee(det_id: str, request: Request, user: dict = Depends(get_current_user)):
    doc = await db_module.db.detainees.find_one({"_id": db_module._oid(det_id)})
    if not doc:
        raise HTTPException(404, "Không tìm thấy hồ sơ")
    _ensure_can_touch(doc, user)
    cid = doc.get("case_id")
    if cid is not None:
        case_doc = await db_module.db.cases.find_one({"_id": cid})
        if case_doc:
            db_module._ensure_case_editable(case_doc)
    await db_module.db.detainees.delete_one({"_id": db_module._oid(det_id)})
    await db_module.db.scene_matches.delete_many({"detainee_id": db_module._oid(det_id)})
    if cid is not None:
        await db_module.db.cases.update_one({"_id": cid}, {"$set": {"updated_at": datetime.utcnow()}})
    await db_module._log(request, user, "delete", "detainee", doc.get("personal_id", det_id), ref_id=det_id, case_id=cid)
    return {"ok": True}


@router.post("/api/detainees/{det_id}/transfer")
async def transfer_detainee(det_id: str, body: TransferBody, request: Request, user: dict = Depends(get_current_user)):
    doc = await db_module.db.detainees.find_one({"_id": db_module._oid(det_id)})
    if not doc:
        raise HTTPException(404, "Không tìm thấy hồ sơ")
    _ensure_can_touch(doc, user)
    new_code = (body.cell_code or "").strip()
    if new_code and not await db_module.db.cells.find_one({"code": new_code}):
        raise HTTPException(400, f"Buồng {new_code} không tồn tại")
    old_code = doc.get("cell_code") or ""
    if old_code == new_code:
        raise HTTPException(400, "Nghi phạm đã ở buồng này")
    await db_module.db.detainees.update_one(
        {"_id": db_module._oid(det_id)},
        {"$set": {"cell_code": new_code or None, "updated_at": datetime.utcnow()}},
    )
    await db_module._log(
        request, user, "update", "detainee", doc.get("personal_id", det_id),
        {"transfer": {"from": old_code, "to": new_code}}, ref_id=det_id,
    )
    return {"ok": True, "from": old_code, "to": new_code}
